import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from app.schemas.models import SyllabusResponse

EXPORTS_DIR = "outputs/exports"


def sf(run, size=11, bold=False, color=None, italic=False, font="Times New Roman"):
    run.font.name   = font
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def remove_borders(table):
    for row in table.rows:
        for cell in row.cells:
            tc  = cell._tc.get_or_add_tcPr()
            bdr = OxmlElement("w:tcBorders")
            for side in ["top","left","bottom","right","insideH","insideV"]:
                b = OxmlElement(f"w:{side}")
                b.set(qn("w:val"),   "none")
                b.set(qn("w:sz"),    "0")
                b.set(qn("w:space"), "0")
                b.set(qn("w:color"), "auto")
                bdr.append(b)
            tc.append(bdr)


def red_footer(doc, programme, branch):
    p   = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    txt = f"  Syllabus of {programme} \u2013 {branch}{'  '*30}PAGE 1"
    r   = p.add_run(txt)
    sf(r, size=9, bold=True, color=(255,255,255), font="Arial")
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "CC0000")
    pPr.append(shd)


def export_syllabus_to_docx(syllabus: SyllabusResponse) -> str:
    doc = Document()

    for sec in doc.sections:
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    # ── Syllabus heading ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Syllabus")
    sf(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(6)

    # ── Course title ──
    title = syllabus.course_name
    if syllabus.course_code:
        title += f" ({syllabus.course_code})"
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(title)
    sf(r2, size=13, bold=True)
    p2.paragraph_format.space_after = Pt(8)

    # ── Programme info ──
    parts = []
    if syllabus.programme:       parts.append(syllabus.programme.upper())
    if syllabus.education_level: parts.append(syllabus.education_level.title())
    if syllabus.year_of_study:   parts.append(f"Year {syllabus.year_of_study}")
    if syllabus.semester:        parts.append(f"Semester-{syllabus.semester}")
    if syllabus.branch:          parts.append(f"Branch: {syllabus.branch}")
    if parts:
        pi = doc.add_paragraph()
        pr = pi.add_run(" | ".join(parts))
        sf(pr, size=10, italic=True, color=(80,80,80))
        pi.paragraph_format.space_after = Pt(8)

    # ── L:T:P + Credits ──
    ltp_tbl = doc.add_table(rows=1, cols=2)
    ltp_tbl.columns[0].width = Inches(3.0)
    ltp_tbl.columns[1].width = Inches(3.5)
    remove_borders(ltp_tbl)
    ll  = ltp_tbl.cell(0,0).paragraphs[0]
    lr  = ltp_tbl.cell(0,1).paragraphs[0]
    lr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    llr = ll.add_run(f"L:T:P:: {syllabus.ltp or '3:1:0'}")
    lrr = lr.add_run(f"Credits-{syllabus.credits or 4}")
    sf(llr, size=11, bold=True)
    sf(lrr, size=11, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ── Course Objectives ──
    if syllabus.course_objectives:
        p = doc.add_paragraph()
        sf(p.add_run("COURSE OBJECTIVES: "), size=11, bold=True)
        sf(p.add_run("The objectives of the course are to:"), size=11)
        p.paragraph_format.space_after = Pt(4)
        for obj in syllabus.course_objectives:
            bp = doc.add_paragraph(style="List Number")
            sf(bp.add_run(obj), size=11)
            bp.paragraph_format.left_indent  = Inches(0.5)
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after  = Pt(2)
        doc.add_paragraph()

    # ── Course Outcomes ──
    if syllabus.course_outcomes:
        p = doc.add_paragraph()
        sf(p.add_run("COURSE OUTCOMES:"), size=11, bold=True)
        p.paragraph_format.space_after = Pt(4)
        p2 = doc.add_paragraph()
        sf(p2.add_run("At the end of this course, the students will be able to:"), size=11)
        p2.paragraph_format.space_after = Pt(4)
        for co in syllabus.course_outcomes:
            bp = doc.add_paragraph(style="List Number")
            sf(bp.add_run(co), size=11)
            bp.paragraph_format.left_indent  = Inches(0.5)
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after  = Pt(2)
        doc.add_paragraph()

    # ── Units ──
    for unit in syllabus.units:
        hours = unit.hours or 8

        up = doc.add_paragraph()
        up.paragraph_format.space_before = Pt(10)
        up.paragraph_format.space_after  = Pt(4)
        sf(up.add_run(f"{unit.unit_id}: {unit.unit_title}:"), size=11, bold=True)
        up.add_run("\t\t\t\t")
        sf(up.add_run(f"({hours} hours)"), size=11, bold=True)

        if unit.topics_paragraph:
            tp = doc.add_paragraph()
            sf(tp.add_run(unit.topics_paragraph), size=11)
            tp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            tp.paragraph_format.space_after = Pt(6)
        elif unit.topics:
            tp = doc.add_paragraph()
            sf(tp.add_run(", ".join(unit.topics) + f". ({hours} hours)"), size=11)
            tp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            tp.paragraph_format.space_after = Pt(6)

        if unit.topics:
            p = doc.add_paragraph()
            sf(p.add_run("Topics Covered:"), size=11, bold=True)
            p.paragraph_format.space_after = Pt(2)
            for topic in unit.topics:
                bp = doc.add_paragraph(style="List Bullet")
                sf(bp.add_run(topic), size=11)
                bp.paragraph_format.left_indent  = Inches(0.5)
                bp.paragraph_format.space_before = Pt(1)
                bp.paragraph_format.space_after  = Pt(1)

        if unit.unit_objectives:
            p = doc.add_paragraph()
            sf(p.add_run("Course Specific Objectives (CSOs):"), size=11, bold=True)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(2)
            for obj in unit.unit_objectives:
                bp = doc.add_paragraph(style="List Bullet")
                sf(bp.add_run(obj), size=11)
                bp.paragraph_format.left_indent  = Inches(0.5)
                bp.paragraph_format.space_before = Pt(1)
                bp.paragraph_format.space_after  = Pt(1)

        if unit.unit_outcomes:
            p = doc.add_paragraph()
            sf(p.add_run("Course Specific Outcomes:"), size=11, bold=True)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(2)
            for out in unit.unit_outcomes:
                bp = doc.add_paragraph(style="List Bullet")
                sf(bp.add_run(out), size=11)
                bp.paragraph_format.left_indent  = Inches(0.5)
                bp.paragraph_format.space_before = Pt(1)
                bp.paragraph_format.space_after  = Pt(1)

        if unit.assessments:
            p = doc.add_paragraph()
            sf(p.add_run("Assessments:"), size=11, bold=True)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(2)
            for a in unit.assessments:
                bp = doc.add_paragraph(style="List Bullet")
                sf(bp.add_run(a), size=11)
                bp.paragraph_format.left_indent  = Inches(0.5)
                bp.paragraph_format.space_before = Pt(1)
                bp.paragraph_format.space_after  = Pt(1)

        if unit.readings:
            p = doc.add_paragraph()
            sf(p.add_run("Readings:"), size=11, bold=True)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(2)
            for rd in unit.readings:
                bp = doc.add_paragraph(style="List Bullet")
                sf(bp.add_run(rd), size=11)
                bp.paragraph_format.left_indent  = Inches(0.5)
                bp.paragraph_format.space_before = Pt(1)
                bp.paragraph_format.space_after  = Pt(1)

        doc.add_paragraph()

    # ── Books ──
    if syllabus.textbooks:
        p = doc.add_paragraph()
        sf(p.add_run("BOOKS:"), size=11, bold=True)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after  = Pt(6)
        for book in syllabus.textbooks:
            bp = doc.add_paragraph(style="List Number")
            sf(bp.add_run(book), size=11)
            bp.paragraph_format.left_indent  = Inches(0.5)
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after  = Pt(2)
        doc.add_paragraph()

    # ── YouTube ──
    if syllabus.youtube_resources:
        p = doc.add_paragraph()
        sf(p.add_run("YOUTUBE & VIDEO RESOURCES:"), size=11, bold=True)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(4)
        for res in syllabus.youtube_resources:
            bp = doc.add_paragraph(style="List Number")
            sf(bp.add_run(res), size=11)
            bp.paragraph_format.left_indent  = Inches(0.5)
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after  = Pt(2)
        doc.add_paragraph()

    # ── Open Source ──
    if syllabus.open_source_resources:
        p = doc.add_paragraph()
        sf(p.add_run("OPEN SOURCE & ONLINE RESOURCES:"), size=11, bold=True)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(4)
        for res in syllabus.open_source_resources:
            bp = doc.add_paragraph(style="List Number")
            sf(bp.add_run(res), size=11)
            bp.paragraph_format.left_indent  = Inches(0.5)
            bp.paragraph_format.space_before = Pt(2)
            bp.paragraph_format.space_after  = Pt(2)

    # ── Red footer ──
    prog = syllabus.programme.upper() if syllabus.programme else "BTECH"
    br   = syllabus.branch or "Computer Science and Engineering"
    red_footer(doc, prog, br)

    # ── Save ──
    os.makedirs(EXPORTS_DIR, exist_ok=True)
    safe  = syllabus.course_name.replace(" ", "_")
    code  = f"_{syllabus.course_code}"       if syllabus.course_code   else ""
    prog  = syllabus.programme.upper()       if syllabus.programme     else "general"
    yr    = f"_Year{syllabus.year_of_study}" if syllabus.year_of_study else ""
    sem   = f"_S{syllabus.semester}"         if syllabus.semester      else ""
    fname = f"{EXPORTS_DIR}/{safe}{code}_{prog}{yr}{sem}_syllabus.docx"
    doc.save(fname)
    print(f"DOCX saved: {fname}")
    return os.path.abspath(fname)